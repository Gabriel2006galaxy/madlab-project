import docx
from docx.shared import Pt

REQUIRED_HEADINGS = ["CERTIFICATE", "ABSTRACT", "INDEX", "LIST OF FIGURES", "ACKNOWLEDGEMENT"]
BODY_TEXT = ("This paragraph contains enough characters to count as a body "
             "candidate for typography inference and validation purposes.")


def build_report_fixture(path, *, wrong_font=False, wrong_spacing=False,
                          drop_heading=None, rename_heading=None):
    d = docx.Document()
    section = d.sections[0]
    section.left_margin = Pt(72)
    section.right_margin = Pt(72)
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    headings = list(REQUIRED_HEADINGS)
    if drop_heading:
        headings = [h for h in headings if h != drop_heading]
    if rename_heading:
        headings = [rename_heading if h == "ACKNOWLEDGEMENT" else h for h in headings]

    for h in headings:
        hp = d.add_paragraph()
        hp.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        hr = hp.add_run(h)
        hr.bold = True
        hr.font.size = Pt(14)

        bp = d.add_paragraph()
        bp.paragraph_format.line_spacing = 1.0 if wrong_spacing else 1.15
        br = bp.add_run(BODY_TEXT)
        br.font.name = "Arial" if wrong_font else "Times New Roman"
        br.font.size = Pt(12)
        br.bold = False

    d.save(path)


def build_style_resolution_fixture(path):
    """3 paragraphs:
    p0: run has explicit font override (Arial) -> EXPLICIT
    p1: uses a custom paragraph style with its own font (Georgia) set,
        run has no direct override -> INHERITED
    p2: plain Normal paragraph, no direct run formatting, relies on
        docDefaults -> DEFAULT (or INHERITED if Normal style itself sets it —
        verified empirically in the accompanying test)
    """
    d = docx.Document()

    p0 = d.add_paragraph()
    r0 = p0.add_run("Explicit font paragraph")
    r0.font.name = "Arial"
    r0.font.size = Pt(11)

    custom_style = d.styles.add_style("CustomBody", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    custom_style.font.name = "Georgia"
    custom_style.font.size = Pt(13)
    p1 = d.add_paragraph(style="CustomBody")
    p1.add_run("Style-inherited font paragraph")

    p2 = d.add_paragraph()
    p2.add_run("Doc-default font paragraph")

    d.save(path)
